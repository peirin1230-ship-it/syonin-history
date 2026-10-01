"""Excel・HTMLビューア・JSON の出力、ターミナル表示。"""
from __future__ import annotations

import base64
import gzip
import json
import re
import sqlite3
from datetime import date
from pathlib import Path

from .pipeline import Paths


def _con(p: Paths) -> sqlite3.Connection:
    if not p.db.exists():
        raise SystemExit("データベースがありません。先に `python -m mdtrack update` を実行してください。")
    con = sqlite3.connect(p.db)
    con.row_factory = sqlite3.Row
    return con


# ---------------------------------------------------------------------------
# MEDIS 突き合わせ（ローカル出力のみ）
# ---------------------------------------------------------------------------
def medis_compare(p: Paths, con: sqlite3.Connection) -> dict:
    """MEDIS 最新スナップショット（と1つ前）から、承認番号ごとの現況と差分を作る。"""
    mp = p.medis / "medis.sqlite"
    if not mp.exists():
        return {}
    m = sqlite3.connect(mp)
    snaps = [r[0] for r in m.execute("SELECT snapshot FROM medis_snapshots ORDER BY snapshot")]
    if not snaps:
        return {}
    latest = snaps[-1]
    prev = snaps[-2] if len(snaps) > 1 else None

    def state(snap):
        st: dict[str, dict] = {}
        for appr, code, name, price, nd, setting, jan, sales, maker in m.execute(
                "SELECT approval_no, receipt_code, material_name, price, notice_date, setting, jan, sales_name, maker "
                "FROM medis_items WHERE snapshot=?", (snap,)):
            if not appr:
                continue
            d = st.setdefault(appr, {"codes": {}, "jans": set(), "sales": sales, "maker": maker})
            d["jans"].add(jan)
            c = d["codes"].setdefault(code or name, {"code": code, "name": name, "price": price,
                                                     "notice_date": nd, "setting": setting, "n_jan": 0})
            c["n_jan"] += 1
        return st

    cur = state(latest)
    old = state(prev) if prev else {}
    listed: dict[str, set] = {}
    for appr, code in con.execute("SELECT DISTINCT approval_no, category_code FROM listing"):
        listed.setdefault(appr, set()).add(code)
    out = {"snapshot": latest, "prev_snapshot": prev, "by_approval": {}}
    for appr in set(cur) | set(old):
        c = cur.get(appr, {"codes": {}, "jans": set()})
        o = old.get(appr, {"codes": {}, "jans": set()})
        added = [v for k, v in c["codes"].items() if k not in o["codes"]] if prev else []
        removed = [v for k, v in o["codes"].items() if k not in c["codes"]] if prev else []
        changed = [{"code": k, "before": o["codes"][k]["price"], "after": v["price"], "name": v["name"]}
                   for k, v in c["codes"].items() if k in o["codes"] and o["codes"][k]["price"] != v["price"]]
        in_notice = listed.get(appr, set())
        not_in_notice = [v for v in c["codes"].values()
                         if re.fullmatch(r"\d{9}", v["code"] or "") and v["code"] not in in_notice] if in_notice else []
        out["by_approval"][appr] = {
            "codes": list(c["codes"].values()), "n_jan": len(c["jans"]),
            "sales": c.get("sales") or o.get("sales"), "maker": c.get("maker") or o.get("maker"),
            "added": added, "removed": removed, "price_changed": changed,
            "not_in_notice": not_in_notice, "known_in_notice": bool(in_notice),
        }
    return out


# ---------------------------------------------------------------------------
def build_payload(p: Paths, include_medis: bool) -> dict:
    con = _con(p)
    aps = []
    for r in con.execute("SELECT * FROM approvals ORDER BY approval_no"):
        aps.append([r["approval_no"], r["sales_name"] or "", r["applicant"] or "", r["first_date"] or "",
                    r["last_date"] or "", r["kubuns"] or "", r["n_products"], json.loads(r["current_json"]),
                    r["flags"] or ""])
    evs: dict[str, list] = {}
    for r in con.execute("SELECT approval_no, date, type, title, detail, kubun, category_code, price_before, "
                         "price_after, doc_id, source FROM events ORDER BY approval_no, date"):
        evs.setdefault(r[0], []).append(list(r)[1:])
    prods: dict[str, list] = {}
    for r in con.execute("SELECT approval_no, product_code, product_name, MIN(effective_date), "
                         "GROUP_CONCAT(DISTINCT category_no) FROM listing WHERE product_code IS NOT NULL "
                         "GROUP BY approval_no, product_code ORDER BY approval_no, MIN(effective_date)"):
        prods.setdefault(r[0], []).append([r[1], r[2] or "", r[3] or "", r[4] or ""])
    docs = {r["doc_id"]: [r["title"], r["url"], r["notice_date"], r["kind"]]
            for r in con.execute("SELECT doc_id, title, url, notice_date, kind FROM notices")}
    stats = dict(con.execute("SELECT status, COUNT(*) FROM notices GROUP BY status").fetchall())
    span = con.execute("SELECT MIN(effective_date), MAX(effective_date) FROM listing").fetchone()
    built = con.execute("SELECT value FROM meta WHERE key='built_at'").fetchone()
    medis = medis_compare(p, con) if include_medis else {}
    if medis:
        known = {a[0] for a in aps}
        for appr, m in medis["by_approval"].items():
            if appr in known or not m["codes"]:
                continue
            cur = [{"category": c["name"], "kubun": "", "setting": c["setting"], "code": c["code"],
                    "master_name": c["name"], "price_now": c["price"], "price_listed": None, "unit": None,
                    "abolished": None, "since": c.get("notice_date") or "", "n_products": c["n_jan"]}
                   for c in m["codes"]]
            aps.append([appr, m.get("sales") or "", m.get("maker") or "", "", "", "MEDIS", m["n_jan"], cur,
                        "MEDISのみ（通知の解析範囲外）"])
    # 通知ごとの変更点
    nchanges: dict[str, list] = {}
    counts: dict[str, dict] = {}
    for r in con.execute("SELECT doc_id, approval_no, sales_name, applicant, kubun, actions, change_types, summary, "
                         "detail_json FROM notice_changes ORDER BY doc_id, approval_no"):
        d = json.loads(r[8])
        nchanges.setdefault(r[0], []).append([r[1], r[2] or "", r[3] or "", r[4] or "", r[5] or "", r[6], r[7],
                                              d.get("before", []), d.get("after", []), d.get("moves", []),
                                              d.get("n_products", 0), d.get("n_new", 0)])
        c = counts.setdefault(r[0], {})
        for t in r[6].split(","):
            c[t] = c.get(t, 0) + 1
    corr_by_target: dict[str, list] = {}
    for r in con.execute("SELECT DISTINCT target_doc, correction_doc, approval_no FROM corrections "
                         "WHERE target_doc IS NOT NULL"):
        corr_by_target.setdefault(r[0], []).append([r[1], r[2]])
    notices = []
    for r in con.execute("SELECT n.doc_id, n.notice_date, MIN(c.effective_date), n.kind, n.title, n.url, n.status "
                         "FROM notices n JOIN notice_changes c USING(doc_id) GROUP BY n.doc_id "
                         "ORDER BY MIN(c.effective_date) DESC, n.notice_date DESC"):
        notices.append([r[0], r[1] or "", r[2] or "", r[3], r[4], r[5], r[6], counts.get(r[0], {}),
                        len(nchanges.get(r[0], []))])
    payload = {
        "built_at": built[0] if built else "", "span": list(span), "notice_stats": stats,
        "approvals": aps, "events": evs, "products": prods, "docs": docs, "medis": medis,
        "notices": notices, "nchanges": nchanges, "corr_by_target": corr_by_target,
    }
    return payload


def _encode(payload: dict) -> str:
    raw = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return base64.b64encode(gzip.compress(raw, 9)).decode("ascii")


def write_viewer(p: Paths, path: Path, include_medis: bool, fragment: bool = False) -> int:
    """ビューアHTMLを書き出す。fragment=True はドキュメント骨格なし（Artifact 公開用）。"""
    tpl = (Path(__file__).parent / "viewer_template.html").read_text(encoding="utf-8")
    b64 = _encode(build_payload(p, include_medis))
    body = tpl.replace("__DATA_B64__", b64)
    if fragment:
        html = body
    else:
        i = body.index('<header class="top">')
        html = ("<!doctype html>\n<html lang=\"ja\">\n<head>\n<meta charset=\"utf-8\">\n"
                "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1, viewport-fit=cover\">\n"
                "<style>body{margin:0}[hidden]{display:none!important}</style>\n"
                + body[:i] + "</head>\n<body>\n" + body[i:] + "\n</body>\n</html>\n")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(html, encoding="utf-8")
    return len(html)


_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _xl(v):
    """Excel に書けない制御文字（OCR の読取ゴミなど）を除く。"""
    return _ILLEGAL.sub("", v) if isinstance(v, str) else v


def write_excel(p: Paths, path: Path, with_listing: bool = True) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    con = _con(p)
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(color="FFFFFF", bold=True)

    def sheet(title, headers, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append([_xl(v) for v in r])
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    wb.remove(wb.active)
    rows = []
    for r in con.execute("SELECT * FROM approvals ORDER BY approval_no"):
        cur = json.loads(r["current_json"])
        cats = " / ".join(f"{c['master_name'] or c['category']}" for c in cur)
        prices = " / ".join(f"{c['price_now']:,.0f}" if c.get("price_now") is not None else "—" for c in cur)
        rows.append((r["approval_no"], r["sales_name"], r["applicant"], r["first_date"], r["last_date"],
                     r["kubuns"], r["n_products"], r["n_categories"], cats, prices, r["n_events"], r["flags"]))
    sheet("承認番号一覧", ["承認番号", "販売名", "保険適用希望者", "初回適用日", "最終掲載日", "区分", "製品数",
                      "機能区分数", "現在の機能区分", "現在の償還価格(円)", "イベント数", "注意"], rows,
          [19, 34, 26, 11, 11, 8, 7, 7, 60, 22, 8, 24])
    ev = con.execute("SELECT e.approval_no, a.sales_name, e.date, e.title, e.detail, e.kubun, e.category_code, "
                     "e.price_before, e.price_after, e.source, e.doc_id FROM events e "
                     "LEFT JOIN approvals a USING(approval_no) ORDER BY e.date DESC, e.approval_no").fetchall()
    sheet("変更イベント", ["承認番号", "販売名", "日付", "種別", "内容", "区分", "特定器材コード", "変更前(円)",
                      "変更後(円)", "出典", "通知ID"], ev, [19, 30, 11, 22, 70, 7, 12, 11, 11, 10, 14])
    li = [] if not with_listing else con.execute("SELECT l.effective_date, l.approval_no, l.sales_name, l.product_name, l.product_code, "
                     "l.applicant, l.setting, l.kubun, CASE l.action WHEN 'new' THEN '新規' WHEN 'add' THEN "
                     "'追加・変更' ELSE l.action END, l.category, l.category_code, l.price_unit, l.price, n.title, "
                     "l.corrected_by, CASE l.ocr WHEN 1 THEN 'OCR' ELSE '' END, l.category_ocr "
                     "FROM listing l LEFT JOIN notices n USING(doc_id) "
                     "ORDER BY l.effective_date, l.approval_no").fetchall()
    if with_listing:
        sheet("掲載明細", ["適用開始日", "承認番号", "販売名", "製品名", "製品コード", "保険適用希望者", "医科/歯科",
                        "区分", "掲載種別", "決定機能区分", "特定器材コード", "単位", "償還価格(円)", "掲載通知",
                        "訂正通知ID", "読取", "OCRの読取文字（機能区分）"], li,
              [11, 19, 28, 34, 16, 24, 8, 6, 9, 50, 12, 10, 11, 50, 14, 6, 40])
    sh = con.execute("SELECT h.code, h.basic_name, h.name, h.beppyo, h.kubun_no, h.valid_from, h.price, h.unit, "
                     "h.abolish_date FROM ssk_history h WHERE h.code IN (SELECT DISTINCT category_code FROM listing) "
                     "ORDER BY h.code, h.valid_from").fetchall()
    sheet("機能区分の価格履歴", ["特定器材コード", "基本名称", "名称", "別表", "区分番号", "変更年月日", "価格(円)",
                         "単位", "廃止年月日"], sh, [12, 60, 34, 6, 8, 11, 11, 8, 11])
    ncs = con.execute("SELECT c.effective_date, n.notice_date, c.approval_no, c.sales_name, c.applicant, c.kubun, "
                      "c.change_types, c.summary, CASE c.ocr WHEN 1 THEN 'OCR' ELSE '' END, n.title "
                      "FROM notice_changes c LEFT JOIN notices n USING(doc_id) "
                      "ORDER BY c.effective_date DESC, c.approval_no").fetchall()
    sheet("通知ごとの変更点", ["適用開始日", "通知日", "承認番号", "販売名", "保険適用希望者", "区分", "変更の種類",
                         "内容", "読取", "通知"], ncs, [11, 11, 19, 30, 24, 7, 22, 80, 6, 50])
    nt = con.execute("SELECT notice_date, effective_date, CASE kind WHEN 'notice' THEN '通知' WHEN 'correction' "
                     "THEN '訂正' WHEN 'replacement' THEN '差替' ELSE kind END, title, status, n_rows, url "
                     "FROM notices ORDER BY notice_date DESC").fetchall()
    sheet("通知一覧", ["通知日", "適用日", "種別", "標題", "状態", "行数", "URL"], nt, [11, 11, 6, 70, 18, 7, 60])
    if "通知ごとの変更点" in wb.sheetnames:
        wb.move_sheet("通知ごとの変更点", offset=-wb.sheetnames.index("通知ごとの変更点"))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def export_all(p: Paths, log=print) -> None:
    p.out.mkdir(parents=True, exist_ok=True)
    n = write_viewer(p, p.out / "viewer.html", include_medis=True)
    log(f"  → {p.out / 'viewer.html'}（{n / 1e6:.1f}MB、MEDIS突合せ含む・ローカル用）")
    n = write_viewer(p, p.root / "docs" / "index.html", include_medis=False)
    log(f"  → {p.root / 'docs' / 'index.html'}（{n / 1e6:.1f}MB、公的情報のみ・共有用）")
    write_excel(p, p.out / "syonin_history.xlsx")
    log(f"  → {p.out / 'syonin_history.xlsx'}")


def print_history(p: Paths, approval_no: str) -> None:
    con = _con(p)
    a = approval_no.strip().upper()
    r = con.execute("SELECT * FROM approvals WHERE approval_no=?", (a,)).fetchone()
    if not r:
        print(f"{a}: B・C区分の掲載が見つかりません")
        return
    print(f"{a}  {r['sales_name']}  （{r['applicant']}）")
    print(f"  初回適用 {r['first_date']} / 区分 {r['kubuns']} / 製品 {r['n_products']}件")
    for c in json.loads(r["current_json"]):
        pn = f"¥{c['price_now']:,.0f}" if c.get("price_now") is not None else "—"
        ab = f"（{c['abolished']} 廃止）" if c.get("abolished") else ""
        print(f"  現在: {c['master_name'] or c['category']}  {pn} {ab}")
    print("  履歴:")
    for e in con.execute("SELECT date, title, detail FROM events WHERE approval_no=? ORDER BY date", (a,)):
        print(f"   {e[0]}  {e[1]}  {e[2][:100] if e[2] else ''}")
